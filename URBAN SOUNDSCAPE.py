!pip install openpyxl

import numpy as np
import pandas as pd
import os
import warnings
import json
from datetime import datetime
from tqdm import tqdm
import pickle
from collections import Counter
import openpyxl

warnings.filterwarnings('ignore')

import librosa
import librosa.display
import soundfile as sf
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestClassifier, GradientBoostingClassifier
from sklearn.discriminant_analysis import LinearDiscriminantAnalysis
from sklearn.svm import SVC
from sklearn.metrics import accuracy_score, precision_score, recall_score, f1_score, confusion_matrix, classification_report
import matplotlib.pyplot as plt
import seaborn as sns

print("Libraries imported!")

class Config:
    DATASET_PATH = "/content/data/ESC-50-master"
    AUDIO_PATH = "/content/data/ESC-50-master/audio"
    METADATA_PATH = "/content/data/ESC-50-master/meta/esc50.csv"

    SAMPLE_RATE = 22050
    DURATION = 5

    N_MFCC = 40
    N_FFT = 2048
    HOP_LENGTH = 512

    CLASS_LABELS = [
        'dog', 'rooster', 'pig', 'cow', 'frog', 'cat', 'hen', 'insect', 'sheep', 'crow',
        'rain', 'sea', 'crackling_fire', 'crickets', 'chirping_birds', 'water_drops', 'wind', 'pouring_water', 'toilet_flush', 'thunderstorm',
        'crying_baby', 'sneezing', 'clapping', 'breathing', 'coughing', 'footsteps', 'laughing', 'brushing_teeth', 'snoring', 'drinking_sipping',
        'door_wood_knock', 'mouse_click', 'keyboard_typing', 'door_wood_creaks', 'can_opening', 'washing_machine', 'vacuum_cleaner', 'clock_alarm', 'clock_tick', 'glass_breaking',
        'helicopter', 'chainsaw', 'siren', 'car_horn', 'engine', 'train', 'church_bells', 'airplane', 'fireworks', 'hand_saw'
    ]
    NUM_CLASSES = len(CLASS_LABELS)

    MODEL_DIR = "/content/models"
    RESULTS_DIR = "/content/results"
    FIGURES_DIR = "/content/figures"

    @classmethod
    def create_directories(cls):
        for dir_path in [cls.MODEL_DIR, cls.RESULTS_DIR, cls.FIGURES_DIR]:
            os.makedirs(dir_path, exist_ok=True)

Config.create_directories()
print(f"Configuration loaded! {Config.NUM_CLASSES} classes")


import urllib.request
import zipfile

print("Downloading ESC-50 (600MB, 10-15 minutes)...")

os.makedirs('/content/data', exist_ok=True)

if not os.path.exists(Config.DATASET_PATH):
    url = "https://github.com/karolpiczak/ESC-50/archive/master.zip"
    file_path = "/content/data/ESC-50.zip"

    urllib.request.urlretrieve(url, file_path)
    print("Download complete!")

    print("Extracting...")
    with zipfile.ZipFile(file_path, 'r') as zip_ref:
        zip_ref.extractall('/content/data/')
    print("Extracted!")
else:
    print("Dataset already exists!")

print(f"Dataset ready at: {Config.DATASET_PATH}")


class ESC50DataLoader:
    def __init__(self):
        self.metadata = None

    def load_metadata(self):
        print("Loading metadata...")
        self.metadata = pd.read_csv(Config.METADATA_PATH)
        print(f"Loaded {len(self.metadata)} samples, {self.metadata['target'].nunique()} classes")
        return self.metadata

    def load_audio_file(self, file_path):
        try:
            audio, _ = librosa.load(file_path, sr=Config.SAMPLE_RATE, duration=Config.DURATION)
            if len(audio) < Config.SAMPLE_RATE * Config.DURATION:
                audio = np.pad(audio, (0, Config.SAMPLE_RATE * Config.DURATION - len(audio)))
            return audio
        except:
            return None

    def load_dataset(self, folds=None, max_samples_per_class=None):
        if self.metadata is None:
            self.load_metadata()

        df = self.metadata[self.metadata['fold'].isin(folds)] if folds else self.metadata

        if max_samples_per_class:
            df = df.groupby('target').head(max_samples_per_class)

        audio_data = []
        labels = []
        file_names = []

        print(f"Loading {len(df)} audio files...")
        for idx, row in tqdm(df.iterrows(), total=len(df)):
            file_path = os.path.join(Config.AUDIO_PATH, row['filename'])
            if os.path.exists(file_path):
                audio = self.load_audio_file(file_path)
                if audio is not None:
                    audio_data.append(audio)
                    labels.append(row['target'])
                    file_names.append(row['filename'])

        print(f"Loaded {len(audio_data)} samples")
        return np.array(audio_data), np.array(labels), file_names

print("Data loader defined!")

class AudioPreprocessor:
    def __init__(self):
        self.config = Config

    def extract_mfcc(self, audio):
        mfcc = librosa.feature.mfcc(y=audio, sr=Config.SAMPLE_RATE, n_mfcc=Config.N_MFCC,
                                    n_fft=Config.N_FFT, hop_length=Config.HOP_LENGTH)
        return mfcc

    def extract_chroma(self, audio):
        chroma = librosa.feature.chroma_stft(y=audio, sr=Config.SAMPLE_RATE, n_fft=Config.N_FFT,
                                              hop_length=Config.HOP_LENGTH)
        return chroma

    def extract_spectral_features(self, audio):
        spectral_centroids = librosa.feature.spectral_centroid(y=audio, sr=Config.SAMPLE_RATE)
        spectral_rolloff = librosa.feature.spectral_rolloff(y=audio, sr=Config.SAMPLE_RATE)
        zero_crossing_rate = librosa.feature.zero_crossing_rate(audio)
        return spectral_centroids, spectral_rolloff, zero_crossing_rate

    def extract_statistical_features(self, feature_matrix):
        mean_val = np.mean(feature_matrix, axis=1)
        std_val = np.std(feature_matrix, axis=1)
        max_val = np.max(feature_matrix, axis=1)
        min_val = np.min(feature_matrix, axis=1)
        median_val = np.median(feature_matrix, axis=1)
        return np.concatenate([mean_val, std_val, max_val, min_val, median_val])

    def process_audio(self, audio):
        features_list = []

        mfcc = self.extract_mfcc(audio)
        features_list.append(self.extract_statistical_features(mfcc))

        chroma = self.extract_chroma(audio)
        features_list.append(self.extract_statistical_features(chroma))

        spectral_centroids, spectral_rolloff, zcr = self.extract_spectral_features(audio)
        statistical_spectral = np.concatenate([
            [np.mean(spectral_centroids), np.std(spectral_centroids)],
            [np.mean(spectral_rolloff), np.std(spectral_rolloff)],
            [np.mean(zcr), np.std(zcr)]
        ])
        features_list.append(statistical_spectral)

        all_features = np.concatenate(features_list)
        return all_features

    def batch_process(self, audio_list):
        features_list = []
        for audio in tqdm(audio_list, desc="Extracting features"):
            features = self.process_audio(audio)
            features_list.append(features)
        return np.array(features_list)

print("Preprocessor defined!")

# ==================== NEW: SEGMENTATION FOR LDA ====================
def segment_audio(audio, sample_rate=Config.SAMPLE_RATE, window_duration=1.0, hop_duration=0.5):
    """Segment audio into overlapping windows for LDA multi-sound detection"""
    window_length = int(window_duration * sample_rate)
    hop_length = int(hop_duration * sample_rate)
    segments = []
    for start in range(0, len(audio) - window_length + 1, hop_length):
        segment = audio[start:start + window_length]
        segments.append(segment)
    return segments

def extract_features_for_segments(audio_segments, preprocessor):
    """Extract features for each segment"""
    features = [preprocessor.process_audio(seg) for seg in audio_segments]
    return np.array(features)

print("Segmentation functions defined!")
# ===================================================================

print("Loading ESC-50 data...")

loader = ESC50DataLoader()
metadata = loader.load_metadata()

print("\nLoading training data (folds 1-4)...")
X_train_audio, y_train, _ = loader.load_dataset(folds=[1, 2, 3, 4])

print("Loading test data (fold 5)...")
X_test_audio, y_test, _ = loader.load_dataset(folds=[5])

print(f"\nTrain: {len(X_train_audio)} samples")
print(f"Test: {len(X_test_audio)} samples")
print(f"Classes: {Config.NUM_CLASSES}")

print("Extracting audio features...\n")

preprocessor = AudioPreprocessor()

print("Training data:")
X_train = preprocessor.batch_process(X_train_audio)

print("\nTest data:")
X_test = preprocessor.batch_process(X_test_audio)

print(f"\nTrain features shape: {X_train.shape}")
print(f"Test features shape: {X_test.shape}")

# Split training data for validation
X_train_split, X_val, y_train_split, y_val = train_test_split(
    X_train, y_train, test_size=0.2, random_state=42, stratify=y_train
)

# Standardize features
scaler = StandardScaler()
X_train_split = scaler.fit_transform(X_train_split)
X_val = scaler.transform(X_val)
X_test = scaler.transform(X_test)

print(f"\nTrain: {X_train_split.shape}")
print(f"Val: {X_val.shape}")
print(f"Test: {X_test.shape}")

print("\n" + "="*60)
print("Training Machine Learning Models")
print("="*60)

# Train Random Forest Classifier
print("\n1. Training Random Forest Classifier...")
rf_model = RandomForestClassifier(n_estimators=100, random_state=42, n_jobs=-1, verbose=1)
rf_model.fit(X_train_split, y_train_split)
print("Random Forest trained!")

y_pred_rf = rf_model.predict(X_test)
y_pred_proba_rf = rf_model.predict_proba(X_test)

accuracy_rf = accuracy_score(y_test, y_pred_rf)
precision_rf = precision_score(y_test, y_pred_rf, average='weighted', zero_division=0)
recall_rf = recall_score(y_test, y_pred_rf, average='weighted', zero_division=0)
f1_rf = f1_score(y_test, y_pred_rf, average='weighted', zero_division=0)

print("\n" + "="*60)
print("Random Forest Results:")
print("="*60)
print(f"Accuracy:  {accuracy_rf:.4f}")
print(f"Precision: {precision_rf:.4f}")
print(f"Recall:    {recall_rf:.4f}")
print(f"F1-Score:  {f1_rf:.4f}")
print("="*60)

# Train Gradient Boosting Classifier
print("\n2. Training Gradient Boosting Classifier...")
gb_model = GradientBoostingClassifier(n_estimators=100, random_state=42, verbose=1)
gb_model.fit(X_train_split, y_train_split)
print("Gradient Boosting trained!")

y_pred_gb = gb_model.predict(X_test)
y_pred_proba_gb = gb_model.predict_proba(X_test)

accuracy_gb = accuracy_score(y_test, y_pred_gb)
precision_gb = precision_score(y_test, y_pred_gb, average='weighted', zero_division=0)
recall_gb = recall_score(y_test, y_pred_gb, average='weighted', zero_division=0)
f1_gb = f1_score(y_test, y_pred_gb, average='weighted', zero_division=0)

print("\n" + "="*60)
print("Gradient Boosting Results:")
print("="*60)
print(f"Accuracy:  {accuracy_gb:.4f}")
print(f"Precision: {precision_gb:.4f}")
print(f"Recall:    {recall_gb:.4f}")
print(f"F1-Score:  {f1_gb:.4f}")
print("="*60)

# Train SVM Classifier
print("\n3. Training Support Vector Machine (SVM)...")
svm_model = SVC(kernel='rbf', C=100, gamma='scale', probability=True, random_state=42, verbose=1)
svm_model.fit(X_train_split, y_train_split)
print("SVM trained!")

y_pred_svm = svm_model.predict(X_test)
y_pred_proba_svm = svm_model.predict_proba(X_test)

accuracy_svm = accuracy_score(y_test, y_pred_svm)
precision_svm = precision_score(y_test, y_pred_svm, average='weighted', zero_division=0)
recall_svm = recall_score(y_test, y_pred_svm, average='weighted', zero_division=0)
f1_svm = f1_score(y_test, y_pred_svm, average='weighted', zero_division=0)

print("\n" + "="*60)
print("SVM Results:")
print("="*60)
print(f"Accuracy:  {accuracy_svm:.4f}")
print(f"Precision: {precision_svm:.4f}")
print(f"Recall:    {recall_svm:.4f}")
print(f"F1-Score:  {f1_svm:.4f}")
print("="*60)

# ==================== NEW: TRAIN LDA ====================
print("\n4. Training LDA for Multi-Sound Detection...")
segmented_features_list = []
segmented_labels = []

for audio, label in tqdm(zip(X_train_audio, y_train), total=len(X_train_audio), desc="Segmenting audio"):
    segments = segment_audio(audio)
    feats = extract_features_for_segments(segments, preprocessor)
    segmented_features_list.append(feats)
    segmented_labels.extend([label] * len(segments))

X_train_segments = np.vstack(segmented_features_list)
y_train_segments = np.array(segmented_labels)
X_train_segments_scaled = scaler.fit_transform(X_train_segments)

lda_model = LinearDiscriminantAnalysis()
lda_model.fit(X_train_segments_scaled, y_train_segments)
print("LDA Model trained!")
# ========================================================

# Model Comparison
print("\n" + "="*60)
print("Model Comparison:")
print("="*60)
comparison_data = {
    'Model': ['Random Forest', 'Gradient Boosting', 'SVM'],
    'Accuracy': [accuracy_rf, accuracy_gb, accuracy_svm],
    'Precision': [precision_rf, precision_gb, precision_svm],
    'Recall': [recall_rf, recall_gb, recall_svm],
    'F1-Score': [f1_rf, f1_gb, f1_svm]
}
comparison_df = pd.DataFrame(comparison_data)
print(comparison_df.to_string(index=False))
# ==================== EXPORT ANALYTICS TO EXCEL ====================

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

excel_path = "/content/results/model_results_dashboard.xlsx"

# Save dataframe and confusion matrix
with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
    
    comparison_df.to_excel(writer, sheet_name="Model_Comparison", index=False)

    cm_df = pd.DataFrame(cm)
    cm_df.to_excel(writer, sheet_name="Confusion_Matrix", index=False)

print("Excel analytics file created!")

# Load workbook to add charts
wb = load_workbook(excel_path)
ws = wb["Model_Comparison"]

# Accuracy chart
accuracy_chart = BarChart()
accuracy_chart.title = "Model Accuracy Comparison"
accuracy_chart.y_axis.title = "Accuracy"

data = Reference(ws, min_col=2, min_row=1, max_row=4)
cats = Reference(ws, min_col=1, min_row=2, max_row=4)

accuracy_chart.add_data(data, titles_from_data=True)
accuracy_chart.set_categories(cats)

ws.add_chart(accuracy_chart, "H2")

# Metrics chart
metric_chart = BarChart()
metric_chart.title = "Precision Recall F1 Comparison"

data2 = Reference(ws, min_col=3, min_row=1, max_col=5, max_row=4)

metric_chart.add_data(data2, titles_from_data=True)
metric_chart.set_categories(cats)

ws.add_chart(metric_chart, "H20")

wb.save(excel_path)

print("Excel dashboard created!")
print("="*60)

best_model_name = comparison_df.loc[comparison_df['Accuracy'].idxmax(), 'Model']
best_model = [rf_model, gb_model, svm_model][comparison_df['Accuracy'].argmax()]
best_y_pred = [y_pred_rf, y_pred_gb, y_pred_svm][comparison_df['Accuracy'].argmax()]
best_y_pred_proba = [y_pred_proba_rf, y_pred_proba_gb, y_pred_proba_svm][comparison_df['Accuracy'].argmax()]

print(f"\nBest Model: {best_model_name}")

# Confusion Matrix for Best Model
cm = confusion_matrix(y_test, best_y_pred)
print(f"Confusion matrix shape: {cm.shape}")

plt.figure(figsize=(20, 18))
sns.heatmap(cm, annot=True, fmt='d', cmap='Blues', xticklabels=Config.CLASS_LABELS,
            yticklabels=Config.CLASS_LABELS, cbar_kws={'label': 'Count'}, annot_kws={'size': 6})
plt.title(f'ESC-50 Confusion Matrix - {best_model_name} (50 Classes)', fontsize=16, fontweight='bold')
plt.ylabel('True Label')
plt.xlabel('Predicted Label')
plt.xticks(rotation=45, ha='right', fontsize=7)
plt.yticks(rotation=0, fontsize=7)
plt.tight_layout()
plt.savefig('/content/figures/confusion_matrix.png', dpi=100, bbox_inches='tight')
plt.show()
print("Confusion matrix saved!")

# Feature Importance Plot
if best_model_name in ['Random Forest', 'Gradient Boosting']:
    feature_importance = best_model.feature_importances_
    top_indices = np.argsort(feature_importance)[-15:]

    plt.figure(figsize=(10, 6))
    plt.barh(range(len(top_indices)), feature_importance[top_indices])
    plt.yticks(range(len(top_indices)), [f"Feature {i}" for i in top_indices])
    plt.xlabel('Importance')
    plt.title(f'Top 15 Feature Importances - {best_model_name}', fontweight='bold')
    plt.tight_layout()
    plt.savefig('/content/figures/feature_importance.png', dpi=100, bbox_inches='tight')
    plt.show()
    print("Feature importance plot saved!")

# Model Comparison Visualization
fig, axes = plt.subplots(2, 2, figsize=(14, 10))

metrics = ['Accuracy', 'Precision', 'Recall', 'F1-Score']
for idx, metric in enumerate(metrics):
    ax = axes[idx // 2, idx % 2]
    values = comparison_df[metric].values
    bars = ax.bar(comparison_df['Model'], values, color=['#1f77b4', '#ff7f0e', '#2ca02c'])
    ax.set_ylabel(metric)
    ax.set_title(f'{metric} Comparison', fontweight='bold')
    ax.set_ylim([0, 1])
    for bar, val in zip(bars, values):
        height = bar.get_height()
        ax.text(bar.get_x() + bar.get_width()/2., height,
                f'{val:.4f}', ha='center', va='bottom', fontsize=10)
    ax.grid(axis='y', alpha=0.3)

plt.tight_layout()
plt.savefig('/content/figures/model_comparison.png', dpi=100, bbox_inches='tight')
plt.show()
print("Model comparison plot saved!")

# ==================== NEW: SAVE LDA MODEL ====================
with open('/content/models/random_forest_model.pkl', 'wb') as f:
    pickle.dump(rf_model, f)
with open('/content/models/gradient_boosting_model.pkl', 'wb') as f:
    pickle.dump(gb_model, f)
with open('/content/models/svm_model.pkl', 'wb') as f:
    pickle.dump(svm_model, f)
with open('/content/models/lda_model.pkl', 'wb') as f:
    pickle.dump(lda_model, f)
with open('/content/models/scaler.pkl', 'wb') as f:
    pickle.dump(scaler, f)
with open('/content/models/preprocessor.pkl', 'wb') as f:
    pickle.dump(preprocessor, f)
# ===========================================================

print("\nAll models saved!")

# Save results
results = {
    'Random Forest': {'accuracy': float(accuracy_rf), 'precision': float(precision_rf), 'recall': float(recall_rf), 'f1_score': float(f1_rf)},
    'Gradient Boosting': {'accuracy': float(accuracy_gb), 'precision': float(precision_gb), 'recall': float(recall_gb), 'f1_score': float(f1_gb)},
    'SVM': {'accuracy': float(accuracy_svm), 'precision': float(precision_svm), 'recall': float(recall_svm), 'f1_score': float(f1_svm)},
    'Best Model': best_model_name
}

with open('/content/results/results.json', 'w') as f:
    json.dump(results, f, indent=4)

print("Results saved!")

# Prediction function for custom audio
from google.colab import files

print("\n" + "="*60)
print("Upload your audio file for prediction:")
print("="*60)
uploaded = files.upload()

audio_file = list(uploaded.keys())[0]
audio_path = f"/content/{audio_file}"

print(f"\nUploaded: {audio_file}")

def predict(audio_path):
    """Single sound prediction using best model"""
    audio, sr = librosa.load(audio_path, sr=Config.SAMPLE_RATE, duration=Config.DURATION)
    if len(audio) < Config.SAMPLE_RATE * Config.DURATION:
        audio = np.pad(audio, (0, Config.SAMPLE_RATE * Config.DURATION - len(audio)))

    features = preprocessor.process_audio(audio)
    features_scaled = scaler.transform([features])

    pred_proba = best_model.predict_proba(features_scaled)
    pred_class_id = np.argmax(pred_proba[0])
    confidence = pred_proba[0][pred_class_id]
    pred_class = Config.CLASS_LABELS[pred_class_id]

    print("\n" + "="*60)
    print(" SINGLE SOUND PREDICTION")
    print("="*60)
    print(f"Predicted: {pred_class.upper()}")
    print(f"Confidence: {confidence:.2%}")

    top_5 = np.argsort(pred_proba[0])[::-1][:5]
    print("\nTop 5:")
    for i, idx in enumerate(top_5, 1):
        print(f"  {i}. {Config.CLASS_LABELS[idx]:20s} {pred_proba[0][idx]:.2%}")

def predict_multi_sound(audio_path):
    """Multi-sound detection using LDA"""
    audio, sr = librosa.load(audio_path, sr=Config.SAMPLE_RATE, duration=Config.DURATION)
    if len(audio) < Config.SAMPLE_RATE * Config.DURATION:
        audio = np.pad(audio, (0, Config.SAMPLE_RATE * Config.DURATION - len(audio)))

    segments = segment_audio(audio)
    segment_features = extract_features_for_segments(segments, preprocessor)
    segment_features_scaled = scaler.transform(segment_features)
    segment_preds = lda_model.predict(segment_features_scaled)

    segment_labels = [Config.CLASS_LABELS[p] for p in segment_preds]
    counts = Counter(segment_labels)

    print("\n" + "="*60)
    print(" MULTI-SOUND DETECTION (LDA)")
    print("="*60)
    print(f"Total segments analyzed: {len(segments)}")
    print(f"Unique sounds detected: {len(counts)}")
    print("\nDetected Sounds:")
    for i, (sound, count) in enumerate(counts.most_common(10), 1):
        percentage = (count / len(segments)) * 100
        bar = "█" * int(percentage / 2)
        print(f"  {i}. {sound:20s} {bar:25s} {count:3d} ({percentage:5.1f}%)")

# Run both predictions
print("\n" + "="*80)
print("RUNNING PREDICTIONS...")
print("="*80)

predict(audio_path)

predict_multi_sound(audio_path)

print("Results saved!")
from google.colab import files
files.download("/content/results/model_results_dashboard.xlsx")
